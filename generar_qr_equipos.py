# -*- coding: utf-8 -*-
"""
Generador de QR limpios para SCV-PTBG
=====================================

Genera un QR NUEVO y LIMPIO por cada equipo, cuyo contenido es EXACTAMENTE el
valor de MOT_EQUIPO.CODE_QR (el TAC del equipo) tal como esta en PTBG_DAT.

Por que asi: el backend (Api_scv_ptbg.py -> get_equipo_by_qr) busca el equipo con
    WHERE UPPER(TRIM(COALESCE(e.CODE_QR,''))) = %s  OR TAGNAME = %s  OR LOCALIZACION = %s
Si cada QR codifica el CODE_QR verbatim, el escaneo hace match DIRECTO y sin las
heuristicas fragiles de hoy (PTBG-004, MOT-6241, numeros sueltos, URLs...).

Fuente de datos: la API REST /equipos (no requiere credenciales de BD ni token).
La API conserva el campo crudo CODE_QR ademas del QR_CODE con fallback, por lo
que aqui podemos distinguir y OMITIR los equipos sin CODE_QR real.

Salida:
  1) qr_equipos_control.csv   -> tabla de control con el contenido exacto de cada QR
  2) qr_omitidos_sin_codeqr.csv -> equipos SIN CODE_QR (para completar en la BD)
  3) qr_png/*.png             -> imagen QR por equipo (si 'qrcode' esta instalado)

Uso:
  python generar_qr_equipos.py
  python generar_qr_equipos.py --api http://192.168.100.241:8001 --out ./qr_out
  python generar_qr_equipos.py --sin-imagenes        # solo los CSV
  python generar_qr_equipos.py --con-etiqueta        # PNG con texto (equipo/TAC) debajo

Requisitos:
  pip install requests
  pip install "qrcode[pil]"      # solo para generar las imagenes PNG
"""

import argparse
import csv
import os
import re
import sys

try:
    import requests
except ImportError:
    print("ERROR: falta el paquete 'requests'. Instalalo con:  pip install requests")
    sys.exit(1)

DEFAULT_API = "http://192.168.100.241:8001"
VACIOS = {"", "NULL", "NONE", "N/A", "NAN"}


# ── Normalizacion identica a la del backend (Api_scv_ptbg._normalize_qr) ──────
# La replicamos para DETECTAR codigos que el backend transformaria al escanear,
# ya que esos NO harian match exacto contra CODE_QR y son "riesgosos".
def normalize_qr_backend(value):
    if value is None:
        return ""
    code = str(value).strip()
    if "/" in code:
        code = code.rstrip("/").split("/")[-1]
    if "code=" in code.lower():
        code = code.split("=")[-1]
    code = code.strip().upper()
    m = re.fullmatch(r"PTBG[-_ ]*(\d+)", code)
    if m:
        return str(int(m.group(1)))
    m = re.fullmatch(r"LOC[-_ ]*(\d+)", code)
    if m:
        return str(int(m.group(1)))
    return code


def es_vacio(v):
    return v is None or str(v).strip().upper() in VACIOS


def limpiar(v):
    return "" if v is None else str(v).strip()


def slug_archivo(texto):
    """Nombre de archivo seguro a partir del CODE_QR."""
    s = re.sub(r"[^A-Za-z0-9._-]+", "_", str(texto).strip())
    return s.strip("_") or "SIN_CODIGO"


def obtener_equipos(api_base):
    url = api_base.rstrip("/") + "/equipos"
    print(f"-> Consultando {url} ...")
    try:
        r = requests.get(url, timeout=30)
    except requests.RequestException as e:
        print(f"ERROR: no se pudo conectar a la API ({e}).")
        print("  Verifica que el backend este corriendo y que esta maquina alcance esa IP.")
        sys.exit(1)
    if r.status_code != 200:
        print(f"ERROR: la API respondio HTTP {r.status_code}: {r.text[:300]}")
        sys.exit(1)
    data = r.json()
    if not isinstance(data, list):
        print("ERROR: /equipos no devolvio una lista JSON.")
        sys.exit(1)
    print(f"  {len(data)} equipos recibidos.")
    return data


def generar_png(contenido, ruta, etiqueta=None):
    """Genera un PNG del QR. Devuelve True si lo creo, False si falta la libreria."""
    try:
        import qrcode
        from qrcode.constants import ERROR_CORRECT_M
    except ImportError:
        return False

    qr = qrcode.QRCode(
        version=None,                 # tamano automatico segun el contenido
        error_correction=ERROR_CORRECT_M,
        box_size=10,
        border=4,
    )
    qr.add_data(contenido)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white").convert("RGB")

    if etiqueta:
        try:
            from PIL import Image, ImageDraw, ImageFont
            texto = etiqueta.strip()
            font = ImageFont.load_default()
            draw = ImageDraw.Draw(img)
            # medir texto
            try:
                bbox = draw.textbbox((0, 0), texto, font=font)
                tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
            except Exception:
                tw, th = draw.textsize(texto, font=font)
            pad = 12
            lienzo = Image.new("RGB", (img.width, img.height + th + pad * 2), "white")
            lienzo.paste(img, (0, 0))
            d2 = ImageDraw.Draw(lienzo)
            d2.text(((img.width - tw) / 2, img.height + pad), texto, fill="black", font=font)
            img = lienzo
        except Exception:
            pass  # si algo falla con la fuente, guardamos el QR sin texto

    img.save(ruta)
    return True


def _qr_pil(contenido, target_px=None, border=4):
    """QR como PIL.Image, NITIDO y con zona de silencio correcta.

    - border=4 modulos: zona de silencio que exige el estandar (ML Kit/zxing la
      necesitan; sin ella la app no detecta el codigo).
    - Si se pide target_px, se escoge un box_size ENTERO para acercarse a ese
      tamano SIN interpolacion, de modo que los modulos queden con bordes duros
      (un resize bicubico los emborrona y rompe el escaneo).
    """
    import qrcode
    from qrcode.constants import ERROR_CORRECT_M
    qr = qrcode.QRCode(error_correction=ERROR_CORRECT_M, border=border, box_size=1)
    qr.add_data(contenido)
    qr.make(fit=True)
    total_mod = qr.modules_count + 2 * border
    qr.box_size = max(4, target_px // total_mod) if target_px else 10
    return qr.make_image(fill_color="black", back_color="white").convert("RGB")


def generar_excel(validos, ruta_xlsx):
    """Catalogo Excel: una fila por equipo con su QR embebido. Devuelve True/False."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.utils import get_column_letter
        from io import BytesIO
    except ImportError:
        return False

    wb = Workbook()
    ws = wb.active
    ws.title = "QR Equipos"

    cabeceras = ["LC", "EQUIPO", "SISTEMA", "TAGNAME", "TAC (CODE_QR)", "QR", "ADVERTENCIA"]
    anchos = [6, 30, 26, 16, 20, 24, 34]
    for i, (h, w) in enumerate(zip(cabeceras, anchos), start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", start_color="0B3D62")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"
    borde = Border(*(Side(style="thin", color="D0D7DE"),) * 4)

    QR_PX = 120  # tamano objetivo del QR embebido, en pixeles
    fila = 2
    for e in validos:
        vals = [e["LOCALIZACION"], e["EQUIPO"], e["SISTEMA"], e["TAGNAME"],
                e["CODE_QR"], "", e["ADVERTENCIA"]]
        for col, v in enumerate(vals, start=1):
            c = ws.cell(row=fila, column=col, value=v)
            c.font = Font(name="Arial", size=10)
            c.border = borde
            c.alignment = Alignment(
                vertical="center",
                horizontal="center" if col in (1, 5) else "left",
                wrap_text=(col in (2, 3, 7)),
            )
        if e["ADVERTENCIA"]:
            ws.cell(row=fila, column=7).font = Font(name="Arial", size=10, color="B00020", bold=True)

        # QR embebido en la columna F (nitido, a tamano real, sin reescalar)
        img_pil = _qr_pil(e["QR_CONTENT"], target_px=QR_PX)
        buf = BytesIO()
        img_pil.save(buf, format="PNG")
        buf.seek(0)
        ximg = XLImage(buf)
        ximg.width, ximg.height = img_pil.size
        ws.add_image(ximg, f"F{fila}")

        ws.row_dimensions[fila].height = img_pil.size[1] * 0.75 + 6  # px -> puntos
        fila += 1

    ws.auto_filter.ref = f"A1:G{fila - 1}"
    wb.save(ruta_xlsx)
    return True


def generar_pdf(validos, ruta_pdf, cols=3, rows=4):
    """Hoja imprimible A4 con los QR en cuadricula (QR + TAC + equipo). True/False."""
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        return False

    PW, PH = 1240, 1754          # A4 aprox a 150 dpi
    margin = 45
    cw = (PW - 2 * margin) // cols
    ch = (PH - 2 * margin) // rows

    def load(sz, bold=False):
        nombres = (["arialbd.ttf", "arial.ttf"] if bold else ["arial.ttf"])
        for n in nombres:
            try:
                return ImageFont.truetype("C:/Windows/Fonts/" + n, sz)
            except Exception:
                pass
        return ImageFont.load_default()

    f_tac = load(30, bold=True)
    f_name = load(22)

    def recorta(draw, texto, font, max_w):
        if draw.textbbox((0, 0), texto, font=font)[2] <= max_w:
            return texto
        while texto and draw.textbbox((0, 0), texto + "...", font=font)[2] > max_w:
            texto = texto[:-1]
        return texto + "..."

    pages = []
    per = cols * rows
    for i in range(0, len(validos), per):
        page = Image.new("RGB", (PW, PH), "white")
        d = ImageDraw.Draw(page)
        for j, e in enumerate(validos[i:i + per]):
            r, c = divmod(j, cols)
            x0 = margin + c * cw
            y0 = margin + r * ch
            d.rectangle([x0 + 3, y0 + 3, x0 + cw - 3, y0 + ch - 3],
                        outline="#C8C8C8", width=1)

            qs = min(cw, ch) - 90
            qr = _qr_pil(e["QR_CONTENT"], target_px=qs)   # nitido, sin resize
            iw = qr.size[0]
            qx = x0 + (cw - iw) // 2
            qy = y0 + 12 + (qs - iw) // 2
            page.paste(qr, (qx, qy))

            ty = y0 + 12 + qs + 2
            tac = recorta(d, str(e["CODE_QR"]), f_tac, cw - 20)
            nom = recorta(d, f"LC{e['LOCALIZACION']}  {e['EQUIPO']}", f_name, cw - 20)
            for txt, fnt, yy in ((tac, f_tac, ty), (nom, f_name, ty + 34)):
                w = d.textbbox((0, 0), txt, font=fnt)[2]
                d.text((x0 + (cw - w) // 2, yy), txt, font=fnt, fill="black")
        pages.append(page)

    if not pages:
        return False
    pages[0].save(ruta_pdf, save_all=True, append_images=pages[1:], resolution=150.0)
    return True


def main():
    ap = argparse.ArgumentParser(description="Generador de QR limpios (CODE_QR) para SCV-PTBG")
    ap.add_argument("--api", default=DEFAULT_API, help=f"URL base de la API (def: {DEFAULT_API})")
    ap.add_argument("--out", default="qr_salida", help="Carpeta de salida (def: qr_salida)")
    ap.add_argument("--sin-imagenes", action="store_true", help="Genera solo los CSV, sin PNG")
    ap.add_argument("--con-etiqueta", action="store_true", help="Agrega texto (equipo/TAC) debajo del QR")
    ap.add_argument("--excel", action="store_true", help="Genera tambien un catalogo Excel (.xlsx) con los QR embebidos")
    ap.add_argument("--pdf", action="store_true", help="Genera una hoja PDF imprimible con los QR en cuadricula")
    args = ap.parse_args()

    # La consola de Windows suele usar cp1252; forzamos UTF-8 para acentos/ñ.
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

    os.makedirs(args.out, exist_ok=True)
    dir_png = os.path.join(args.out, "qr_png")
    if not args.sin_imagenes:
        os.makedirs(dir_png, exist_ok=True)

    equipos = obtener_equipos(args.api)

    validos = []       # equipos con CODE_QR real
    omitidos = []      # equipos sin CODE_QR
    vistos = {}        # CODE_QR (upper) -> lista de equipos (para detectar duplicados)

    for e in equipos:
        code_qr = limpiar(e.get("CODE_QR"))   # campo CRUDO de MOT_EQUIPO
        loc = limpiar(e.get("LOCALIZACION"))
        nombre = limpiar(e.get("EQUIPO"))
        tag = limpiar(e.get("TAGNAME"))
        idv = limpiar(e.get("ID"))
        sistema = limpiar(e.get("SISTEMA"))

        if es_vacio(code_qr):
            omitidos.append({
                "ID": idv, "EQUIPO": nombre, "LOCALIZACION": loc,
                "TAGNAME": tag, "SISTEMA": sistema,
                "MOTIVO": "CODE_QR vacio en MOT_EQUIPO",
            })
            continue

        # Chequeo de seguridad: el backend normaliza al escanear. Si el resultado
        # difiere del CODE_QR en mayusculas, ese codigo NO haria match exacto.
        norm = normalize_qr_backend(code_qr)
        riesgo = "" if norm == code_qr.upper() else f"El backend lo normaliza a '{norm}'"

        vistos.setdefault(code_qr.upper(), []).append(nombre or loc or idv)

        validos.append({
            "ID": idv,
            "EQUIPO": nombre,
            "LOCALIZACION": loc,
            "SISTEMA": sistema,
            "TAGNAME": tag,
            "CODE_QR": code_qr,
            "QR_CONTENT": code_qr,   # <-- contenido EXACTO que se codifica en el QR
            "ADVERTENCIA": riesgo,
        })

    # marcar duplicados (dos equipos con el mismo CODE_QR -> escaneo ambiguo)
    for fila in validos:
        dup = vistos.get(fila["CODE_QR"].upper(), [])
        if len(dup) > 1:
            extra = "CODE_QR DUPLICADO (compartido por %d equipos)" % len(dup)
            fila["ADVERTENCIA"] = (fila["ADVERTENCIA"] + "; " + extra).strip("; ")

    # ── CSV de control ───────────────────────────────────────────────────────
    ruta_ctrl = os.path.join(args.out, "qr_equipos_control.csv")
    with open(ruta_ctrl, "w", newline="", encoding="utf-8-sig") as f:
        cols = ["ID", "EQUIPO", "LOCALIZACION", "SISTEMA", "TAGNAME",
                "CODE_QR", "QR_CONTENT", "ADVERTENCIA", "ARCHIVO_PNG"]
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for fila in validos:
            fila["ARCHIVO_PNG"] = ("" if args.sin_imagenes
                                   else f"LC{fila['LOCALIZACION'] or '000'}_{slug_archivo(fila['CODE_QR'])}.png")
            w.writerow(fila)

    # ── CSV de omitidos ──────────────────────────────────────────────────────
    ruta_omit = os.path.join(args.out, "qr_omitidos_sin_codeqr.csv")
    with open(ruta_omit, "w", newline="", encoding="utf-8-sig") as f:
        cols = ["ID", "EQUIPO", "LOCALIZACION", "SISTEMA", "TAGNAME", "MOTIVO"]
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(omitidos)

    # ── PNG por equipo ───────────────────────────────────────────────────────
    png_ok = 0
    lib_falta = False
    if not args.sin_imagenes:
        for fila in validos:
            etiqueta = None
            if args.con_etiqueta:
                etiqueta = f"{fila['CODE_QR']}  |  LC{fila['LOCALIZACION']}"
            ruta = os.path.join(dir_png, fila["ARCHIVO_PNG"])
            creado = generar_png(fila["QR_CONTENT"], ruta, etiqueta)
            if not creado:
                lib_falta = True
                break
            png_ok += 1

    # ── Excel ────────────────────────────────────────────────────────────────
    ruta_xlsx = os.path.join(args.out, "qr_equipos_catalogo.xlsx")
    excel_ok = False
    if args.excel:
        excel_ok = generar_excel(validos, ruta_xlsx)

    ruta_pdf = os.path.join(args.out, "qr_equipos_imprimir.pdf")
    pdf_ok = False
    if args.pdf:
        pdf_ok = generar_pdf(validos, ruta_pdf)

    # ── Resumen ──────────────────────────────────────────────────────────────
    print("\n--------------------------------------------")
    print("  RESUMEN")
    print("--------------------------------------------")
    print(f"  Equipos con CODE_QR (QR generados): {len(validos)}")
    print(f"  Equipos SIN CODE_QR (omitidos)    : {len(omitidos)}")
    riesgos = [v for v in validos if v['ADVERTENCIA']]
    if riesgos:
        print(f"  [!] Codigos con advertencia       : {len(riesgos)}  (revisa la columna ADVERTENCIA)")
    print(f"\n  CSV control  : {ruta_ctrl}")
    print(f"  CSV omitidos : {ruta_omit}")
    if args.excel:
        if excel_ok:
            print(f"  Excel catalogo: {ruta_xlsx}")
        else:
            print("  [!] No se genero Excel: falta 'openpyxl'.  pip install openpyxl")
    if args.pdf:
        if pdf_ok:
            print(f"  PDF imprimible: {ruta_pdf}")
        else:
            print("  [!] No se genero PDF: falta Pillow.  pip install \"qrcode[pil]\"")
    if not args.sin_imagenes:
        if lib_falta:
            print("\n  [!] No se generaron PNG: falta la libreria 'qrcode'.")
            print("    Instalala con:  pip install \"qrcode[pil]\"   y vuelve a ejecutar.")
            print("    (Los CSV si se generaron correctamente.)")
        else:
            print(f"  PNG generados: {png_ok}  en  {dir_png}")
    print("--------------------------------------------")
    if omitidos:
        print(f"\n  NOTA: {len(omitidos)} equipos no tienen CODE_QR en la BD y fueron omitidos.")
        print("        Completa MOT_EQUIPO.CODE_QR para esos equipos y vuelve a ejecutar.")


if __name__ == "__main__":
    main()
