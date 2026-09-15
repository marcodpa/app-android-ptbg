from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
CAPTURES = ROOT / "capturas_app"
OUTPUT = ROOT / "Informe_visual_STER_2026-08-04.docx"

NAVY = RGBColor(10, 28, 62)
BLUE = RGBColor(31, 77, 120)
TEAL = RGBColor(0, 155, 130)
MUTED = RGBColor(87, 101, 121)
LIGHT = "E8EEF5"
PALE = "F4F6F9"
RED = RGBColor(155, 28, 28)
GOLD = RGBColor(122, 90, 0)


SCREENS = [
    ("01_inicio/01_inicio_modo_noche.png", "Inicio - modo noche",
     "Panel principal industrial con usuario activo, estado USB, accesos directos, indicadores diarios y ruta de trabajo.",
     "La jerarquía visual es clara y el contraste general es adecuado. El contador de equipos del inicio puede permanecer en cero aunque el módulo Equipos tenga datos locales; conviene revisar la actualización de este resumen."),
    ("01_inicio/02_inicio_modo_dia.png", "Inicio - modo día",
     "Variante clara del mismo panel. Mantiene tarjetas azul marino y usa fondo blanco como superficie principal.",
     "El cambio de tema funciona. La estructura y el sidebar se mantienen estables; los módulos conservan su identidad cromática."),
    ("02_equipos/01_modulo_equipos.png", "Módulo Equipos",
     "Listado de activos con búsqueda, filtros por sistema, imagen del conjunto, identificación, localización y accesos a reporte, impresión y servicios.",
     "La pantalla presenta correctamente las configuraciones Motor-Caja-Bomba y ventiladores. Los tres equipos de demostración se utilizaron solo para esta revisión visual."),
    ("03_qr/01_escanear_qr.png", "Escáner QR",
     "Entrada rápida al flujo de trabajo mediante cámara. Incluye guía visual, activación explícita y alternativa manual.",
     "La pantalla es comprensible y evita abrir la cámara sin una acción del operador."),
    ("03_qr/02_ingreso_codigo_manual.png", "Ingreso manual de código",
     "Despliega un campo de identificación y búsqueda como respaldo cuando el QR no puede leerse.",
     "La alternativa manual queda integrada sin abandonar el módulo QR."),
    ("04_servicios/01_seleccion_servicios.png", "Selección de servicios",
     "Permite escoger uno o varios trabajos para un mismo equipo y mantenerlos asociados a una ODT común.",
     "Incluye Vibración, Temperatura, Lubricación, Alineación, Reemplazo y Cambio de coupling según la configuración del equipo."),
    ("04_servicios/02_vibracion_captura.png", "Servicio de Vibración",
     "Captura guiada punto por punto en orientaciones Horizontal, Vertical y Axial. Presenta progreso, ubicación gráfica, lectura anterior, unidad mm/s y estado USB.",
     "Hallazgo visual: la tarjeta blanca de datos técnicos usa texto demasiado claro; Modelo, Serial y Marca pierden contraste. El resto de la captura es legible y operativo."),
    ("04_servicios/03_temperatura_captura.png", "Servicio de Temperatura",
     "Captura una lectura en grados Celsius por punto, con progreso, referencia gráfica, lectura anterior y alerta inmediata de nivel mientras se escribe.",
     "El campo activo ahora usa turquesa y se lee correctamente en modo noche. Persiste el bajo contraste de los datos técnicos sobre la tarjeta blanca."),
    ("04_servicios/04_lubricacion_captura.png", "Servicio de Lubricación",
     "Muestra la ficha del lubricante, el conjunto mecánico y la ubicación exacta de aplicación. Para cada punto informa referencia eléctrica/manual y cantidad cuando existe en la base.",
     "La imagen del equipo y el bloque COLOCAR AQUÍ orientan bien al mecánico. Las cantidades ausentes pueden permanecer sin valor predeterminado, según la regla definida."),
    ("04_servicios/05_alineacion_captura.png", "Servicio de Alineación",
     "Presenta una referencia coaxial del conjunto y organiza ángulos y compensaciones verticales/horizontales para Motor-Caja y Caja-Bomba.",
     "Hallazgo importante: el título ALINEACIÓN MOTOR-CAJA aparece en azul muy oscuro sobre fondo oscuro y casi no se distingue. Debe usar texto primario claro o turquesa."),
    ("04_servicios/06_reemplazo_equipo.png", "Servicio de Reemplazo",
     "Permite seleccionar Motor, Caja y/o Bomba, ingresar la nueva placa y decidir si se actualizan las especificaciones técnicas del motor.",
     "La pantalla inicial es clara y conserva la imagen del conjunto. El botón de revisión permanece deshabilitado hasta seleccionar y completar un componente."),
    ("04_servicios/07_cambio_coupling.png", "Servicio de Cambio de coupling",
     "Registro simple Sí/No para equipos compatibles, con observaciones opcionales y datos del conjunto.",
     "Las dos opciones tienen diferenciación semántica clara: rojo para No y turquesa para Sí."),
    ("05_sincronizacion/01_sincronizacion_usb.png", "Sincronización USB",
     "Centraliza descarga, trabajo offline y envío por USB. Resume pendientes, sincronizados y errores, y muestra acciones de edición/eliminación por tipo de registro.",
     "Se comprobó que Temperatura ya presenta su icono rojo de eliminación, igual que los demás servicios pendientes."),
    ("06_reportes/01_historial_mediciones.png", "Historial de mediciones",
     "Consulta datos MariaDB y Tablet por servicio. En Vibración muestra H/V/A, RMS, fecha, equipo y observaciones.",
     "La estructura facilita comparar registros y distinguir el último dato disponible por equipo."),
    ("06_reportes/02_generar_reporte_equipo.png", "Generar reporte PDF",
     "Selector de reporte integral o específico: Vibración, Temperatura, Alineación, Lubricación, Reemplazos, Coupling, Correas y Ficha técnica.",
     "Hallazgo crítico de contraste: en modo día, el panel inferior conserva fondo azul oscuro pero usa texto negro, haciendo títulos y descripciones difíciles de leer."),
    ("06_reportes/03_imprimir_formatos_equipo.png", "Imprimir formatos del equipo",
     "Permite elegir una de las últimas tres ODT y seleccionar los servicios incluidos antes de imprimir el formato oficial.",
     "Hallazgo crítico de contraste: el panel inferior también muestra texto negro sobre azul oscuro. Debe heredar texto claro cuando la superficie sea oscura."),
    ("07_configuracion/01_ajustes.png", "Ajustes",
     "Presenta sesión activa, política de sincronización exclusiva por USB, versión, tecnologías, unidades y referencia ISO 10816.",
     "La pantalla comunica correctamente que la aplicación no depende de Wi-Fi para sincronizar."),
]


def set_font(run, size=11, color=None, bold=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_title(doc, text, size=26, after=6, color=NAVY):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    set_font(r, size=size, color=color, bold=True)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, bold=True)
        text = text[len(bold_prefix):]
    r = p.add_run(text)
    set_font(r, color=RGBColor(35, 43, 55))
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_font(r)
    return p


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, NAVY, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.text = "STER | Sistema de Trazabilidad de Equipos Rotativos"
    set_font(header.runs[0], size=9, color=MUTED, bold=True)
    add_page_number(section.footer.paragraphs[0])

    # Portada editorial.
    doc.add_paragraph().paragraph_format.space_after = Pt(70)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kr = kicker.add_run("INFORME DE REVISIÓN VISUAL Y FUNCIONAL")
    set_font(kr, size=11, color=TEAL, bold=True)
    title = add_title(doc, "STER", size=34, after=8)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("Sistema de Trazabilidad de Equipos Rotativos")
    set_font(sr, size=17, color=BLUE, bold=True)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(20)
    mr = meta.add_run("Levantamiento completo de pantallas en tablet\n4 de agosto de 2026")
    set_font(mr, size=11, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(55)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nr = note.add_run("17 capturas verificadas · modos noche y día · flujo USB")
    set_font(nr, size=11, color=TEAL, bold=True)

    doc.add_page_break()
    add_title(doc, "Resumen ejecutivo", size=24, color=NAVY)
    add_body(doc, "Se realizó un recorrido visual completo de la aplicación Android STER instalada en la tablet conectada. El levantamiento cubrió navegación principal, equipos, QR, seis servicios, sincronización USB, historial, generación de reportes, impresión y ajustes.")
    add_body(doc, "Para acceder a todas las configuraciones se utilizó una copia local de demostración con tres equipos. Al terminar, la base original de la tablet fue restaurada. No se enviaron datos de demostración a MariaDB.")
    add_heading(doc, "Cobertura", 1)
    for item in (
        "Inicio en modo noche y modo día.",
        "Módulo Equipos y selección múltiple de servicios.",
        "Vibración, Temperatura, Lubricación, Alineación, Reemplazo y Cambio de coupling.",
        "Escaneo QR e ingreso manual.",
        "Sincronización exclusiva por USB y eliminación de pendientes.",
        "Historial, reporte PDF, impresión de formatos y Ajustes.",
    ):
        add_bullet(doc, item)
    add_heading(doc, "Hallazgos prioritarios", 1)
    for item in (
        "Crítico: los paneles de Reporte PDF e Imprimir usan texto negro sobre fondo azul oscuro cuando la app está en modo día.",
        "Alto: el título de la sección Alineación Motor-Caja usa azul oscuro sobre fondo oscuro.",
        "Medio: Modelo, Serial y Marca tienen contraste insuficiente dentro de la tarjeta blanca de Vibración y Temperatura.",
        "Medio: el resumen del inicio puede indicar cero equipos aunque el módulo Equipos tenga activos locales disponibles.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "Matriz de módulos", 1)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(1.8), Inches(3.6), Inches(1.1)]
    for idx, (text, width) in enumerate(zip(("Módulo", "Función", "Estado"), widths)):
        cell = table.rows[0].cells[idx]
        cell.width = width
        cell.text = text
        shade(cell, LIGHT)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_font(cell.paragraphs[0].runs[0], bold=True, color=NAVY)
    set_repeat_table_header(table.rows[0])
    rows = [
        ("Inicio", "Panel, indicadores, USB y rutas", "Revisado"),
        ("Equipos", "Activos, filtros, reportes y servicios", "Revisado"),
        ("Servicios", "Seis flujos operativos", "Revisado"),
        ("Sincronización", "Descarga y subida por USB", "Revisado"),
        ("Reportes", "Historial, PDF e impresión", "Con hallazgo"),
        ("Ajustes", "Sesión, versión e ISO 10816", "Revisado"),
    ]
    for row_data in rows:
        cells = table.add_row().cells
        for idx, (text, width) in enumerate(zip(row_data, widths)):
            cells[idx].width = width
            cells[idx].text = text
            set_cell_margins(cells[idx])
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_font(cells[idx].paragraphs[0].runs[0], size=10)

    for number, (relative, title_text, description, review) in enumerate(SCREENS, start=1):
        doc.add_page_break()
        add_heading(doc, f"{number}. {title_text}", 1)
        image_path = CAPTURES / relative
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_after = Pt(4)
        p_img.add_run().add_picture(str(image_path), width=Inches(3.55))
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(8)
        cr = caption.add_run(f"Figura {number}. {title_text}")
        set_font(cr, size=9, color=MUTED, bold=True)
        add_body(doc, f"Función. {description}", bold_prefix="Función. ")
        observation = add_body(doc, f"Revisión. {review}", bold_prefix="Revisión. ")
        if "Hallazgo crítico" in review:
            for run in observation.runs:
                run.font.color.rgb = RED
        elif "Hallazgo importante" in review:
            for run in observation.runs:
                run.font.color.rgb = GOLD

    doc.add_page_break()
    add_title(doc, "Conclusiones y acciones recomendadas", size=23, color=NAVY)
    add_body(doc, "La aplicación cubre de forma coherente el ciclo operativo: identificar el equipo, seleccionar servicios, capturar lecturas guiadas, conservarlas offline, sincronizar por USB y consultar o imprimir resultados.")
    add_heading(doc, "Acciones de interfaz", 1)
    for item in (
        "Aplicar colores de texto dependientes de la superficie en todos los bottom sheets de reporte e impresión.",
        "Cambiar el encabezado Alineación Motor-Caja a AppColors.textPrimary o AppColors.teal.",
        "Oscurecer el texto técnico mostrado sobre tarjetas blancas en Vibración y Temperatura.",
        "Unificar la recarga del contador del inicio con la lista local de Equipos después de abrir o sincronizar la base.",
    ):
        add_bullet(doc, item)
    add_heading(doc, "Resultado del levantamiento", 1)
    add_body(doc, "Las 17 capturas originales quedan organizadas en la carpeta capturas_app por módulo. El informe documenta el estado observado en la tablet y puede utilizarse como referencia para la siguiente ronda de correcciones y pruebas de regresión.")

    doc.core_properties.title = "Informe visual STER"
    doc.core_properties.subject = "Revisión completa de pantallas y servicios"
    doc.core_properties.author = "Equipo STER"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
