from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
CAP = ROOT / "documentation" / "capturas_informe_corto"
SERVICE_CAP = ROOT / "documentation" / "capturas_app" / "04_servicios"
OUT = ROOT / "output" / "docx" / "Informe_resumido_STER_2026-08-04.docx"

NAVY = "07172F"
BLUE = "10233E"
CYAN = "20C8D8"
GREEN = "23BE91"
TEXT = "17243A"
MUTED = "52647E"
LIGHT = "EEF3F8"
BORDER = "CAD6E3"
WHITE = "FFFFFF"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, size=9, bold=False, color=TEXT):
    run.font.name = "Segoe UI"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Segoe UI")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Segoe UI")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)


def add_text(container, text, size=9, bold=False, color=TEXT, align=WD_ALIGN_PARAGRAPH.LEFT, after=3):
    p = container.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    set_font(p.add_run(text), size, bold, color)
    return p


def add_picture(container, path, width):
    p = container.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(path), width=Inches(width))
    return p


def configure_section(section):
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    section.header_distance = Inches(0.22)
    section.footer_distance = Inches(0.22)


def add_header_footer(section):
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run("STER"), 10, True, NAVY)
    r = p.add_run("    Sistema de Trazabilidad de Equipos Rotativos")
    set_font(r, 7.5, False, MUTED)
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run("Informe resumido de funcionamiento - modo noche"), 7, False, MUTED)


def page_start(doc, title, subtitle):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(title), 19, True, NAVY)
    add_text(doc, subtitle, 9, False, MUTED, after=9)


def next_page(doc):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section)
    add_header_footer(section)


def add_card(cell, image_path, title, body, image_width=2.55):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    set_cell_fill(cell, WHITE)
    set_cell_border(cell)
    set_cell_margins(cell, 100, 110, 110, 110)
    cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
    add_picture(cell, image_path, image_width)
    add_text(cell, title, 10.5, True, NAVY, after=2)
    add_text(cell, body, 8.2, False, MUTED, after=0)


def two_cards(doc, cards):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for cell in table.rows[0].cells:
        cell.width = Inches(3.35)
    for cell, spec in zip(table.rows[0].cells, cards):
        add_card(cell, *spec)
    return table


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_section(doc.sections[0])
    add_header_footer(doc.sections[0])
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Segoe UI"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(3)

    page_start(doc, "STER: operacion clara en campo", "Vista resumida de las areas principales y de los documentos que genera la aplicacion.")
    add_picture(doc, CAP / "01_inicio.png", 4.8)
    add_text(doc, "El inicio concentra al usuario, estado USB, accesos principales y ruta de trabajo. La navegacion lateral mantiene disponibles Inicio, Equipos, Sincronizacion, Historial, tema y Configuracion.", 9.5, False, TEXT, after=0)

    next_page(doc)
    page_start(doc, "Equipos y ruta de trabajo", "La informacion se organiza alrededor de cada equipo rotativo.")
    two_cards(doc, [
        (CAP / "02_equipos.png", "Equipos", "Permite buscar, filtrar y abrir la ficha de cada activo para consultar datos, registrar servicios o generar documentos.", 2.3),
        (CAP / "05_servicios.png", "Servicios", "Al seleccionar un equipo se eligen los trabajos de una misma ODT: vibracion, temperatura, alineacion, lubricacion, coupling, correas o reemplazo.", 2.3),
    ])

    next_page(doc)
    page_start(doc, "Capturas de cada servicio", "Pantallas utilizadas por el mecanico durante el registro en campo.")
    services = [
        ("02_vibracion_captura.png", "Vibracion"),
        ("03_temperatura_captura.png", "Temperatura"),
        ("04_lubricacion_captura.png", "Lubricacion"),
        ("05_alineacion_captura.png", "Alineacion"),
        ("06_reemplazo_equipo.png", "Reemplazo de equipo"),
        ("07_cambio_coupling.png", "Cambio de coupling"),
    ]
    grid = doc.add_table(rows=2, cols=3)
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.autofit = False
    for idx, (name, label) in enumerate(services):
        cell = grid.cell(idx // 3, idx % 3)
        cell.width = Inches(2.18)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        set_cell_border(cell)
        set_cell_margins(cell, 70, 70, 70, 70)
        add_picture(cell, SERVICE_CAP / name, 1.85)
        add_text(cell, label, 8.3, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER, 0)

    next_page(doc)
    page_start(doc, "Identificacion y trazabilidad", "Captura rapida en campo y seguimiento centralizado.")
    two_cards(doc, [
        (CAP / "06_qr.png", "Escanear QR", "Identifica el equipo por su codigo y abre directamente su ficha, reduciendo errores de seleccion.", 2.3),
        (CAP / "08_historial.png", "Historial", "Consulta las ordenes y servicios ya registrados para revisar la evolucion y trazabilidad del activo.", 2.3),
    ])

    next_page(doc)
    page_start(doc, "Trabajo sin conexion WiFi", "La transferencia de datos se realiza por USB entre la tablet y la laptop.")
    two_cards(doc, [
        (CAP / "07_sincronizacion.png", "Sincronizacion USB", "Muestra los registros pendientes, permite revisar la cola y transferir mediciones de forma controlada.", 2.3),
        (CAP / "09_ajustes.png", "Configuracion", "Administra datos del usuario, preferencias, tema visual y parametros necesarios para trabajar en campo.", 2.3),
    ])

    next_page(doc)
    page_start(doc, "Reporte PDF", "Documento tecnico de analisis y comparacion.")
    two_cards(doc, [
        (CAP / "03_reporte_seleccion.png", "1. Seleccion", "Desde el equipo se eligen los sistemas y mediciones que formaran el reporte comparativo.", 2.35),
        (CAP / "10_salida_reporte_integral.png", "2. Resultado", "El PDF presenta datos historicos, tablas comparativas y graficas para interpretar la evolucion del equipo.", 2.65),
    ])
    add_text(doc, "El boton Reporte conserva el formato analitico existente y permite revisar tendencias antes de tomar decisiones de mantenimiento.", 9, False, TEXT, after=0)

    next_page(doc)
    page_start(doc, "Imprimir formato oficial", "Planilla operativa lista para archivo o impresion.")
    two_cards(doc, [
        (CAP / "04_imprimir_seleccion.png", "1. Seleccion", "Se escoge una ODT y uno o varios servicios. Los campos sin informacion se representan con '/' para evitar espacios ambiguos.", 2.35),
        (CAP / "11_salida_formato_impresion.png", "2. Resultado", "STER llena el formato correspondiente al tipo de equipo y lo prepara para abrir e imprimir desde la laptop.", 2.65),
    ])
    add_text(doc, "Reporte e Imprimir son funciones distintas: Reporte analiza y compara; Imprimir genera la planilla oficial de la orden seleccionada.", 9, True, NAVY, after=0)

    next_page(doc)
    page_start(doc, "Resumen de servicios", "Cada registro queda asociado al equipo y a una ODT comun.")
    rows = [
        ("Vibracion", "Lecturas por punto, estado inmediato y evolucion historica."),
        ("Temperatura", "Temperaturas por componente con aviso visual durante la captura."),
        ("Alineacion", "Datos de alineacion del conjunto motor-caja-bomba."),
        ("Lubricacion", "Puntos graficos, gramos o emboladas y tipo de grasera manual/electrica."),
        ("Coupling y correas", "Confirmacion de cambio de coupling y registro del ajuste de correas."),
        ("Reemplazo", "Actualiza primero los datos del equipo para que los demas servicios usen la nueva placa."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.55)
    table.columns[1].width = Inches(5.05)
    hdr = table.rows[0].cells
    for cell, label in zip(hdr, ("Servicio", "Funcion")):
        set_cell_fill(cell, NAVY)
        set_cell_border(cell, NAVY)
        set_cell_margins(cell, 100, 110, 100, 110)
        cell.text = ""
        add_text(cell, label, 9, True, WHITE, after=0)
    set_repeat_table_header(table.rows[0])
    for service, function in rows:
        cells = table.add_row().cells
        for cell in cells:
            set_cell_fill(cell, LIGHT)
            set_cell_border(cell)
            set_cell_margins(cell, 120, 110, 120, 110)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cells[0].text = ""
        cells[1].text = ""
        add_text(cells[0], service, 8.8, True, NAVY, after=0)
        add_text(cells[1], function, 8.8, False, TEXT, after=0)
    add_text(doc, "Resultado: una aplicacion enfocada en el trabajo de campo, con captura guiada, trazabilidad por ODT, sincronizacion USB y salidas documentales profesionales.", 10.5, True, NAVY, after=0)

    doc.core_properties.title = "Informe resumido STER"
    doc.core_properties.subject = "Sistema de Trazabilidad de Equipos Rotativos"
    doc.core_properties.author = "STER"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
