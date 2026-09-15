from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import letter
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "tmp" / "xlsx" / "formatos_rev6_final"
OUTPUT = ROOT / "formatos_pdf_nuevos"


def color_hex(value, default=None):
    if not value or value.type != "rgb" or not value.rgb:
        return default
    raw = value.rgb[-6:]
    return "#" + raw


def row_height(sheet, row):
    return float(sheet.row_dimensions[row].height or 15)


def col_width(sheet, col):
    width = sheet.column_dimensions[get_column_letter(col)].width or 8.43
    return float(width) * 7.0


def wrap_text(text, font_name, font_size, max_width):
    lines = []
    for original in str(text).splitlines() or [""]:
        words = original.split()
        if not words:
            lines.append("")
            continue
        current = words[0]
        for word in words[1:]:
            trial = current + " " + word
            if stringWidth(trial, font_name, font_size) <= max_width:
                current = trial
            else:
                lines.append(current)
                current = word
        lines.append(current)
    return lines


def draw_cell(c, sheet, cell, x, y, w, h):
    fill = color_hex(cell.fill.fgColor)
    if fill and fill.upper() not in {"#000000", "#FFFFFF"}:
        c.setFillColor(fill)
        c.rect(x, y, w, h, fill=1, stroke=0)
    elif cell.fill.fill_type == "solid" and fill:
        c.setFillColor(fill)
        c.rect(x, y, w, h, fill=1, stroke=0)

    border_color = "#000000"
    c.setStrokeColor(border_color)
    for side, coords in (
        (cell.border.left, (x, y, x, y + h)),
        (cell.border.right, (x + w, y, x + w, y + h)),
        (cell.border.top, (x, y + h, x + w, y + h)),
        (cell.border.bottom, (x, y, x + w, y)),
    ):
        if side.style:
            c.setLineWidth(1.0 if side.style in {"medium", "thick", "double"} else 0.45)
            c.line(*coords)

    if cell.value is None or cell.value == "":
        return
    text = str(cell.value)
    bold = bool(cell.font.bold)
    italic = bool(cell.font.italic)
    font_name = "Helvetica"
    if bold and italic:
        font_name = "Helvetica-BoldOblique"
    elif bold:
        font_name = "Helvetica-Bold"
    elif italic:
        font_name = "Helvetica-Oblique"
    font_size = min(float(cell.font.sz or 9), 11.0)
    font_color = color_hex(cell.font.color, "#000000")
    c.setFillColor(font_color)
    c.setFont(font_name, font_size)
    padding = 2.2
    align = cell.alignment.horizontal or "left"
    vertical = cell.alignment.vertical or "center"
    rotation = int(cell.alignment.textRotation or 0)
    if rotation in {90, 180}:
        c.saveState()
        c.translate(x + w / 2, y + h / 2)
        c.rotate(rotation)
        c.drawCentredString(0, -font_size / 3, text)
        c.restoreState()
        return
    lines = wrap_text(text, font_name, font_size, max(3, w - 2 * padding))
    leading = font_size * 1.08
    total = len(lines) * leading
    if vertical == "top":
        baseline = y + h - padding - font_size
    elif vertical == "bottom":
        baseline = y + padding + total - leading
    else:
        baseline = y + (h + total) / 2 - leading
    for index, line in enumerate(lines):
        yy = baseline - index * leading
        if align in {"center", "centerContinuous"}:
            c.drawCentredString(x + w / 2, yy, line)
        elif align == "right":
            c.drawRightString(x + w - padding, yy, line)
        else:
            c.drawString(x + padding, yy, line)


def convert(source, target):
    workbook = load_workbook(source, data_only=True)
    sheet = workbook.active
    page_w, page_h = letter
    margin_x, margin_y = 18.0, 14.0
    raw_widths = [col_width(sheet, col) for col in range(1, sheet.max_column + 1)]
    raw_heights = [row_height(sheet, row) for row in range(1, sheet.max_row + 1)]
    scale = min((page_w - 2 * margin_x) / sum(raw_widths), (page_h - 2 * margin_y) / sum(raw_heights))
    widths = [v * scale for v in raw_widths]
    heights = [v * scale for v in raw_heights]
    left = (page_w - sum(widths)) / 2
    top = page_h - (page_h - sum(heights)) / 2
    xs = [left]
    for value in widths:
        xs.append(xs[-1] + value)
    ys = [top]
    for value in heights:
        ys.append(ys[-1] - value)

    merged_by_anchor = {}
    covered = set()
    for merged in sheet.merged_cells.ranges:
        min_col, min_row, max_col, max_row = merged.bounds
        merged_by_anchor[(min_row, min_col)] = (min_row, min_col, max_row, max_col)
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                if (row, col) != (min_row, min_col):
                    covered.add((row, col))

    c = canvas.Canvas(str(target), pagesize=letter, pageCompression=1)
    c.setTitle(source.stem)
    for row in range(1, sheet.max_row + 1):
        for col in range(1, sheet.max_column + 1):
            if (row, col) in covered:
                continue
            cell = sheet.cell(row, col)
            if (row, col) in merged_by_anchor:
                min_row, min_col, max_row, max_col = merged_by_anchor[(row, col)]
                x = xs[min_col - 1]
                y = ys[max_row]
                w = xs[max_col] - xs[min_col - 1]
                h = ys[min_row - 1] - ys[max_row]
            else:
                x = xs[col - 1]
                y = ys[row]
                w = widths[col - 1]
                h = heights[row - 1]
            draw_cell(c, sheet, cell, x, y, w, h)

    for image in sheet._images:
        try:
            anchor = image.anchor
            col = anchor._from.col
            row = anchor._from.row
            width = image.width * 0.75 * scale
            height = image.height * 0.75 * scale
            x = xs[col]
            y = ys[row] - height
            c.drawImage(ImageReader(BytesIO(image._data())), x, y, width, height, preserveAspectRatio=True, mask="auto")
        except Exception:
            pass
    c.showPage()
    c.save()


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for source in sorted(SOURCE.glob("*.xlsx")):
        target = OUTPUT / f"{source.stem}.pdf"
        convert(source, target)
        print(target)


if __name__ == "__main__":
    main()
